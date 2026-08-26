"""只读 DOCX 基础检查；不联网、不写回原文件。"""
from __future__ import annotations

import json
import sys
from pathlib import Path


def inspect(path: str) -> dict[str, object]:
    from docx import Document

    document = Document(str(Path(path).resolve()))
    paragraphs = [
        {"index": index, "text": paragraph.text, "style": paragraph.style.name if paragraph.style else ""}
        for index, paragraph in enumerate(document.paragraphs, 1)
        if paragraph.text.strip()
    ]
    fonts: dict[str, int] = {}
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            name = run.font.name or "未明确设置"
            fonts[name] = fonts.get(name, 0) + len(run.text or "")
    return {
        "source_file": Path(path).name,
        "paragraph_count": len(paragraphs),
        "table_count": len(document.tables),
        "font_usage": fonts,
        "paragraphs": paragraphs,
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: review_document.py FILE")
    print(json.dumps(inspect(sys.argv[1]), ensure_ascii=False))
